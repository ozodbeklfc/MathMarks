"""Read the coloured Excel gradebook and write it back in the same layout."""
import io
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select

import services
from db import ClassTopic, Student, Topic

FILL_TO_COLOR = {"00B050": "green", "FFFF00": "yellow", "FF0000": "red"}
COLOR_TO_FILL = {v: k for k, v in FILL_TO_COLOR.items()}
PART_RE = re.compile(r"^(.*?)\s*\((toq|juft)\)$", re.I)


def _cell_color(cell):
    fill = cell.fill
    if not fill or fill.fill_type != "solid" or fill.fgColor.type != "rgb":
        return None
    return FILL_TO_COLOR.get(str(fill.fgColor.rgb)[-6:].upper())


CHAPTER_MARK = "Bob"  # A1 of a sheet that has a chapter row above the topics


def _chapter_by_column(ws) -> dict:
    """Row 1 of our own export: chapter names, merged over their topics' columns."""
    out = {}
    for rng in ws.merged_cells.ranges:
        if rng.min_row == 1 == rng.max_row:
            title = ws.cell(1, rng.min_col).value
            for col in range(rng.min_col, rng.max_col + 1):
                out[col] = title
    for col in range(3, ws.max_column + 1):
        if col not in out and ws.cell(1, col).value:
            out[col] = ws.cell(1, col).value
    return {c: services.clean(t) for c, t in out.items() if t and services.clean(t)}


def import_workbook(s, data: bytes, chapter_id=None) -> list[dict]:
    """Add the file's data to the journal. Nothing that is not in the file is deleted.

    Each sheet is a class (matched by name, ignoring spaces and case). Students are matched
    by name; new ones are added. Topics are attached to the class; for every topic in the
    file its cells overwrite the stored marks, including empty cells.

    chapter_id given: every topic in the file belongs to that chapter (one file per chapter).
    chapter_id None: chapters come from the file's own chapter row (our export), else none.

    Two layouts are read: the original gradebook (topic headers in row 1) and our own
    export (chapter names in row 1, topic headers in row 2, "Bob" written in A1).
    """
    wb = load_workbook(io.BytesIO(data))
    report = []
    for ws in wb.worksheets:
        with_chapters = services.clean(ws.cell(1, 1).value or "") == CHAPTER_MARK
        head_row = 2 if with_chapters else 1
        chapter_of = _chapter_by_column(ws) if with_chapters and chapter_id is None else {}
        # Topic headers. "(toq)" / "(juft)" columns of one title form one topic.
        columns = []  # (column index, title, part)
        for col in range(3, ws.max_column + 1):
            head = ws.cell(head_row, col).value
            if head is None or not services.clean(head):
                continue  # the old hand-typed total column has no header
            head = services.clean(head)
            if head.lower() in ("jami", "o'rin"):
                continue  # total and place columns written by our own export
            m = PART_RE.match(head)
            columns.append((col, m.group(1).strip(), m.group(2).lower()) if m else (col, head, "single"))
        rows = [r for r in range(head_row + 1, ws.max_row + 1)
                if isinstance(ws.cell(r, 2).value, str) and services.clean(ws.cell(r, 2).value)]
        if not columns or not rows:
            continue

        cls = services.add_class(s, ws.title)
        in_class = set(s.scalars(select(ClassTopic.topic_id).where(ClassTopic.class_id == cls.id)))
        pos = s.scalar(select(func.max(ClassTopic.position)).where(ClassTopic.class_id == cls.id)) or 0

        topic_ids, chapters_seen = {}, set()
        for col, title, part in columns:
            has_parts = part != "single"
            key = (title, has_parts)
            if key in topic_ids:
                continue
            ch = chapter_id
            if ch is None and col in chapter_of:
                ch = services.add_chapter(s, chapter_of[col]).id
            if ch is not None:
                chapters_seen.add(ch)
            # The same topic name may exist in several chapters: look it up inside its chapter.
            same_chapter = Topic.chapter_id.is_(None) if ch is None else Topic.chapter_id == ch
            t = s.scalar(select(Topic).where(Topic.title == title, Topic.has_parts == has_parts, same_chapter))
            if not t:
                t = services.add_topic(s, title, has_parts, class_ids=[], chapter_id=ch)
            if t.id not in in_class:
                pos += 1
                s.add(ClassTopic(class_id=cls.id, topic_id=t.id, position=pos))
                in_class.add(t.id)
            topic_ids[key] = t.id
        s.flush()

        students = {services.clean(st.full_name).lower(): st
                    for st in s.scalars(select(Student).where(Student.class_id == cls.id))}
        st_pos = max([st.position for st in students.values()] or [0])
        n_marks = n_new = 0
        for r in rows:
            name = services.clean(ws.cell(r, 2).value)
            st = students.get(name.lower())
            if not st:
                st_pos += 1
                st = Student(class_id=cls.id, full_name=name, position=st_pos)
                s.add(st)
                s.flush()
                students[name.lower()] = st
                n_new += 1
            for col, title, part in columns:
                color = _cell_color(ws.cell(r, col))
                services.set_mark(s, st.id, topic_ids[(title, part != "single")], part, color)
                n_marks += bool(color)
        report.append({"class": cls.name, "students": len(rows), "new_students": n_new,
                       "topics": len(topic_ids), "marks": n_marks, "chapters": len(chapters_seen)})
    return report


def export_workbook(s) -> bytes:
    """One sheet per class: chapters in row 1, topics in row 2, students from row 3."""
    wb = Workbook()
    wb.remove(wb.active)
    head_align = Alignment(wrap_text=True, horizontal="center", vertical="center")
    for c in services.list_classes(s):
        d = services.class_data(s, c.id)
        chapter_title = {ch["id"]: ch["title"] for ch in d["chapters"]}
        ws = wb.create_sheet(re.sub(r"[\[\]:*?/\\]", "-", d["name"])[:31])
        ws.cell(1, 1, CHAPTER_MARK).font = Font(bold=True)
        ws.cell(2, 1, "№")
        ws.cell(2, 2, "O'quvchi")
        col, cells, groups = 3, [], []  # groups: [chapter title, first col, last col]
        for t in d["topics"]:
            title = chapter_title.get(t["chapter_id"], "")
            if not groups or groups[-1][0] != title:
                groups.append([title, col, col])
            for part in services.parts_of(t["has_parts"]):
                label = t["title"] if part == "single" else f"{t['title']}\n({part})"
                ws.cell(2, col, label).alignment = head_align
                ws.column_dimensions[get_column_letter(col)].width = 13
                cells.append((col, f"{t['id']}:{part}"))
                groups[-1][2] = col
                col += 1
        for title, first, last in groups:
            if title:
                ws.cell(1, first, title).alignment = head_align
                ws.cell(1, first).font = Font(bold=True)
                if last > first:
                    ws.merge_cells(start_row=1, start_column=first, end_row=1, end_column=last)
        ws.cell(2, col, "Jami").alignment = head_align
        ws.cell(2, col + 1, "O'rin").alignment = head_align
        ws.row_dimensions[1].height = 24
        ws.row_dimensions[2].height = 48
        ws.column_dimensions["A"].width = 5
        ws.column_dimensions["B"].width = 30
        ranked = sorted(d["students"], key=lambda r: (-r["total"], r["name"]))
        for i, st in enumerate(ranked, 1):
            r = i + 2
            ws.cell(r, 1, i)
            ws.cell(r, 2, st["name"])
            for cidx, key in cells:
                color = st["marks"].get(key)
                if color:
                    ws.cell(r, cidx).fill = PatternFill("solid", fgColor="FF" + COLOR_TO_FILL[color])
            ws.cell(r, col, st["total"]).font = Font(bold=True)
            ws.cell(r, col + 1, st["place"])
        ws.freeze_panes = "C3"
    if not wb.worksheets:
        wb.create_sheet("Baholar")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
